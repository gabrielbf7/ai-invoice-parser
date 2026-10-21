"""Tests for the Excel exporter, using fictional data only (no API calls)."""

from openpyxl import load_workbook

from invoice_parser.exporter import export_excel


def test_export_creates_three_sheets_with_rows(tmp_path):
    records = [
        {
            "file": "ok.pdf",
            "data": {
                "vendor_name": "Blue Pine Studio", "invoice_number": "BP-1001",
                "invoice_date": "2026-03-10", "due_date": None, "currency": "USD",
                "line_items": [{"description": "Design", "quantity": 1,
                                "unit_price": 100, "amount": 100}],
                "subtotal": 100, "discount": None, "tax": None, "total": 100,
            },
            "reasons": [],
        },
        {"file": "broken.jpg", "data": None, "reasons": ["extraction failed: timeout"]},
    ]

    path = export_excel(records, tmp_path / "result.xlsx")
    workbook = load_workbook(path)

    assert workbook.sheetnames == ["Invoices", "Line items", "Needs review"]
    assert workbook["Invoices"].max_row == 3        # header + 2 invoices
    assert workbook["Line items"].max_row == 2      # header + 1 item
    assert workbook["Needs review"]["B2"].value == "extraction failed: timeout"
    assert workbook["Invoices"]["L2"].value == "OK"
    assert workbook["Invoices"]["L3"].value == "Needs review"