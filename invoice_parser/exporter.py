"""Write extraction results to an Excel file with three sheets."""

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

INVOICE_FIELDS = ["vendor_name", "invoice_number", "invoice_date", "due_date",
                  "currency", "subtotal", "discount", "tax", "total"]
INVOICE_HEADERS = ["file"] + INVOICE_FIELDS + ["line_items", "status"]
LINE_ITEM_HEADERS = ["file", "invoice_number", "description", "quantity", "unit_price", "amount"]
REVIEW_HEADERS = ["file", "reason"]
DATE_FIELDS = {"invoice_date", "due_date"}


def to_excel_date(value):
    """Turn 'YYYY-MM-DD' into a real date so Excel can sort and filter it."""
    try:
        return date.fromisoformat(value)
    except (TypeError, ValueError):
        return value


def fill_sheet(sheet, headers, rows):
    """Write a bold header, the rows, freeze the header and fit column widths."""
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    for column in sheet.columns:
        longest = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column)
        sheet.column_dimensions[column[0].column_letter].width = min(longest + 2, 50)


def export_excel(records, path):
    """Save records ({'file', 'data', 'reasons'}) to an .xlsx file. Returns the path."""
    invoice_rows, item_rows, review_rows = [], [], []

    for record in records:
        file = record["file"]
        data = record["data"] or {}
        items = data.get("line_items") or []
        status = "Needs review" if record["reasons"] else "OK"

        values = [to_excel_date(data.get(f)) if f in DATE_FIELDS else data.get(f)
                  for f in INVOICE_FIELDS]
        invoice_rows.append([file] + values + [len(items), status])

        for item in items:
            item_rows.append([file, data.get("invoice_number"), item.get("description"),
                              item.get("quantity"), item.get("unit_price"), item.get("amount")])

        for reason in record["reasons"]:
            review_rows.append([file, reason])

    workbook = Workbook()
    invoices = workbook.active
    invoices.title = "Invoices"
    fill_sheet(invoices, INVOICE_HEADERS, invoice_rows)
    fill_sheet(workbook.create_sheet("Line items"), LINE_ITEM_HEADERS, item_rows)
    fill_sheet(workbook.create_sheet("Needs review"), REVIEW_HEADERS, review_rows)

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path